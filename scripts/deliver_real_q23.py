"""Prepare model-derived official cell matrices and verify saved engine-authored Excel."""
import csv
from dataclasses import dataclass
from copy import deepcopy, copy
from datetime import date, datetime, time, timedelta
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import openpyxl
from openpyxl.utils import get_column_letter
from src.q3.reporting import emergency_intervals, time_label

OUT = ROOT/'outputs/q23_candidate_20260911'
ANALYSIS = ROOT/'analysis/real_q23_candidate'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


@dataclass(frozen=True)
class DeliverySpec:
    q4: bool = False

    @property
    def out(self):
        return ROOT / ('outputs/q4_candidate_20260911' if self.q4 else 'outputs/q23_candidate_20260911')

    @property
    def analysis(self):
        return ROOT / ('analysis/real_q4_candidate' if self.q4 else 'analysis/real_q23_candidate')

    def identifier(self, number):
        return f'4-{number}' if self.q4 else str(number)

    def stem(self, number):
        return f'q4_{number}' if self.q4 else f'q{number}'


DEFAULT_SPEC = DeliverySpec()
Q4_SPEC = DeliverySpec(q4=True)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2)+'\n')


def prepare(spec=DEFAULT_SPEC):
    OUT, ANALYSIS = spec.out, spec.analysis
    OUT.mkdir(parents=True, exist_ok=True)
    report = json.loads((ANALYSIS/'result.json').read_text())
    validation = json.loads((ANALYSIS/'independent_validation.json').read_text())
    if not validation['passed'] or validation['result_sha256'] != sha(ANALYSIS/'result.json'):
        raise ValueError('Independent validation missing or stale')
    for number in (2, 3):
        identifier, stem = spec.identifier(number), spec.stem(number)
        template = ROOT/f'CUMCM2026Problems/C题/附件/附件5/result{identifier}.xlsx'
        daily = ANALYSIS/f'{stem}_main.jsonl'
        if sha(daily) != report['strategies'][f'{stem}_main']['daily_file_sha256']:
            raise ValueError('Saved daily result changed')
        rows = [json.loads(line) for line in daily.read_text().splitlines()]
        wb = openpyxl.load_workbook(template)
        if len(rows) != 334 or wb['计划购电量'].max_row != 335:
            raise ValueError('Expected 334 official dates')
        purchases, adjusted, battery, emergency, fees = [], [], [], [], []
        paper = [f'# 问题{identifier}指定日期结果（候选）', '',
            '时间按物理区间列示，单位kWh、元。总费用包含紧急购电及适用的逐次调整费用。', '']
        for i, row in enumerate(rows):
            day = date(2025,2,1)+timedelta(days=i)
            if row['day'] != str(day) or wb['计划购电量'].cell(i+2,1).value.date() != day:
                raise ValueError('Template date mapping mismatch')
            plan, s, costs = row['planned_G'], row['schedule'], row['costs']
            purchases.append(plan+[sum(plan), costs['total_yuan'] if number == 2 else costs['planned_yuan']])
            if number == 3:
                adjusted.append(s['G']+[sum(s['G']), costs['total_yuan']])
            for block in range(6):
                a, b = block*24, (block+1)*24
                battery.append([str(day) if block == 0 else None, f'{4*block}:00-{4*block+4}:00',
                    sum(s['C'][a:b]), sum(s['D'][a:b]), 0 if block == 0 else '24:00' if block == 1 else None,
                    s['E'][0] if block == 0 else s['E'][-1] if block == 1 else None])
            intervals = emergency_intervals(s['U'])
            for k, interval in enumerate(intervals):
                emergency.append([str(day) if k == 0 else None, interval['interval'], interval['purchase_kwh']])
            if not intervals:
                emergency.append([str(day), '无', 0])
            fees.append({'day': str(day), **costs, 'planned_kwh': sum(plan),
                         'final_commitment_kwh': sum(s['G']), 'emergency_kwh': sum(s['U'])})
            if str(day) in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
                paper += [f'## {day}', '', '| 时间段 | 购电量 |', '| --- | ---: |']
                for hour in (10,12,14,16,18,20):
                    paper.append(f'| {hour}:00-{hour}:10 | {s["G"][hour*6]:.6f} |')
                paper += ['', f'原计划购电量 {sum(plan):.6f}；最终承诺购电量 {sum(s["G"]):.6f}；紧急购电量 {sum(s["U"]):.6f}；含紧急购电的实际外网总购电量 {sum(s["G"])+sum(s["U"]):.6f}。',
                    f'原计划费用 {costs["planned_yuan"]:.6f}；调整净费用 {costs["adjustment_yuan"]:.6f}；紧急费用 {costs["emergency_yuan"]:.6f}；总费用 {costs["total_yuan"]:.6f}。', '',
                    '| 时间段 | 充电量 | 放电量 |', '| --- | ---: | ---: |']
                for b in battery[-6:]:
                    paper.append(f'| {b[1]} | {b[2]:.6f} | {b[3]:.6f} |')
                paper += ['', f'0:00储电量 {s["E"][0]:.6f}；24:00储电量 {s["E"][-1]:.6f}。', '',
                          '| 紧急购电时间段 | 购电量 |', '| --- | ---: |']
                paper += [f'| {v["interval"]} | {v["purchase_kwh"]:.6f} |' for v in intervals] or ['| 无 | 0 |']
                paper.append('')
        sheets = {'计划购电量': {'start_column': 1, 'dates': False, 'values': purchases}}
        if number == 3:
            sheets['调整购电量'] = {'start_column': 1, 'dates': False, 'values': adjusted}
        sheets['充放电量'] = {'start_column': 0, 'dates': True, 'values': battery}
        sheets['紧急购电量'] = {'start_column': 0, 'dates': True, 'values': emergency}
        write_json(OUT/f'payload{identifier}.json', {'template_sha256': sha(template), 'daily_file_sha256': sha(daily),
            'scope': 'real-candidate-only', 'semantic_acceptance': False, 'sheets': sheets})
        (ANALYSIS/f'{stem}_paper_tables.md').write_text('\n'.join(paper)+'\n')
        with (ANALYSIS/f'{stem}_daily_fees.csv').open('w', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=list(fees[0])); writer.writeheader(); writer.writerows(fees)
        mapping = [{'slot': i, 'column': get_column_letter(i+2), 'physical_start': time_label(i),
                    'physical_end': time_label(i+1), 'original_template_label': wb['计划购电量'].cell(1,i+2).value,
                    'label_offset_minutes': 10} for i in range(144)]
        with (OUT/f'slot_mapping{identifier}.csv').open('w', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=list(mapping[0])); writer.writeheader(); writer.writerows(mapping)
        wb.close()


def preserve(number, spec=DEFAULT_SPEC):
    """Restore unsupported template font/border metadata around engine-authored cells.

    The generated sheetData is kept; source styles are assigned by original
    address or repeated daily-block role. Other original package parts are copied.
    """
    OUT, ANALYSIS = spec.out, spec.analysis
    identifier, stem = spec.identifier(number), spec.stem(number)
    template = ROOT/f'CUMCM2026Problems/C题/附件/附件5/result{identifier}.xlsx'
    target = OUT/f'result{identifier}_candidate.xlsx'
    payload = json.loads((OUT/f'payload{identifier}.json').read_text())
    if sha(template) != payload['template_sha256']:
        raise ValueError('Template changed')
    # Verify all authored values before compatibility preservation.
    verify(number, check_styles=False, spec=spec)
    original_wb = openpyxl.load_workbook(template)
    with ZipFile(template) as source, ZipFile(target) as authored:
        shared = []
        if 'xl/sharedStrings.xml' in authored.namelist():
            shared = [''.join(si.itertext()) for si in ET.fromstring(authored.read('xl/sharedStrings.xml'))]
        replacements = {}
        for i, name in enumerate(original_wb.sheetnames, 1):
            part = f'xl/worksheets/sheet{i}.xml'
            old, new = ET.fromstring(source.read(part)), ET.fromstring(authored.read(part))
            old_data, new_data = old.find('s:sheetData',NS), deepcopy(new.find('s:sheetData',NS))
            for row in new_data:
                r = int(row.attrib['r'])
                source_row = r if name in ('计划购电量', '调整购电量') or r == 1 else (
                    2+(r-2)%6 if name == '充放电量' else
                    2 if payload['sheets'][name]['values'][r-2][0] is not None else 3)
                for cell in row:
                    address = cell.attrib['r']
                    col = re.match('[A-Z]+', address).group()
                    cell.set('s', str(original_wb[name][f'{col}{source_row}'].style_id))
                    if cell.attrib.get('t') == 's':
                        index = int(cell.find('s:v', NS).text)
                        cell.remove(cell.find('s:v', NS))
                        cell.set('t', 'inlineStr')
                        inline = ET.SubElement(cell, '{'+NS['s']+'}is')
                        ET.SubElement(inline, '{'+NS['s']+'}t').text = shared[index]
            index = list(old).index(old_data)
            old.remove(old_data); old.insert(index, new_data)
            dim = old.find('s:dimension',NS)
            if dim is not None:
                table = payload['sheets'][name]
                end_column = get_column_letter(table['start_column']+len(table['values'][0]))
                dim.set('ref', f'A1:{end_column}{len(table["values"])+1}')
            replacements[part] = ET.tostring(old, encoding='utf-8', xml_declaration=True)
        temporary = target.with_suffix('.tmp.xlsx')
        with ZipFile(temporary, 'w') as final:
            for item in source.infolist():
                final.writestr(item, replacements.get(item.filename, source.read(item.filename)))
    temporary.replace(target)
    original_wb.close()


def verify(number, check_styles=True, spec=DEFAULT_SPEC):
    OUT, ANALYSIS = spec.out, spec.analysis
    identifier, stem = spec.identifier(number), spec.stem(number)
    template = ROOT/f'CUMCM2026Problems/C题/附件/附件5/result{identifier}.xlsx'
    target = OUT/f'result{identifier}_candidate.xlsx'
    payload = json.loads((OUT/f'payload{identifier}.json').read_text())
    if sha(template) != payload['template_sha256'] or sha(ANALYSIS/f'{stem}_main.jsonl') != payload['daily_file_sha256']:
        raise ValueError('Stale template or solved data')
    old, new = openpyxl.load_workbook(template), openpyxl.load_workbook(target)
    if new.sheetnames != old.sheetnames:
        raise ValueError('Worksheet order/name mismatch')
    count = 0
    for name, table in payload['sheets'].items():
        ws = new[name]
        if (ws.max_row, ws.max_column) != (len(table['values'])+1, len(table['values'][0])+table['start_column']):
            raise ValueError('Unexpected dimensions: '+name)
        for row in old[name].iter_rows(min_row=1,max_row=1):
            for cell in row:
                if ws[cell.coordinate].value != cell.value:
                    raise ValueError('Template header changed')
        for i, values in enumerate(table['values'],2):
            if name in ('计划购电量', '调整购电量') and ws.cell(i,1).value != old[name].cell(i,1).value:
                raise ValueError('Template date changed')
            for col, expected in enumerate(values,table['start_column']+1):
                cell = ws.cell(i,col)
                actual = cell.value
                if table['dates'] and col == 1 and expected is not None:
                    if not check_styles and isinstance(actual,(int,float)):
                        actual = openpyxl.utils.datetime.from_excel(actual)
                    if not isinstance(actual,datetime) or str(actual.date()) != expected:
                        raise ValueError(f'Date readback mismatch {name}!{cell.coordinate}')
                elif expected == 0 and isinstance(actual,time):
                    if actual != time(0): raise ValueError('Midnight mismatch')
                elif isinstance(expected,(float,int)):
                    if isinstance(actual,bool) or not isinstance(actual,(float,int)) or not math.isfinite(actual) or abs(actual-expected) > 1e-6:
                        raise ValueError(f'Numeric readback mismatch {name}!{cell.coordinate}: {actual} != {expected}')
                    count += 1
                elif actual != expected:
                    raise ValueError(f'Value mismatch {name}!{cell.coordinate}: {actual} != {expected}')
                if check_styles:
                    r = i if name in ('计划购电量','调整购电量') else 2+(i-2)%6 if name == '充放电量' else 2 if values[0] is not None else 3
                    src = old[name].cell(r,col)
                    if any(copy(getattr(cell,k)) != copy(getattr(src,k)) for k in ('font','fill','border','alignment','protection')) or cell.number_format != src.number_format:
                        raise ValueError('Template style mismatch')
    old.close(); new.close()
    if check_styles:
        with ZipFile(template) as a, ZipFile(target) as b:
            if set(a.namelist()) != set(b.namelist()) or any(a.read(p) != b.read(p) for p in a.namelist() if not p.startswith('xl/worksheets/sheet')):
                raise ValueError('Nonworksheet package changed')
    return {'number': identifier if spec.q4 else number, 'numeric_cells_verified': count, 'all_values_verified': True,
            'template_styles_preserved': check_styles, 'semantic_acceptance': False, 'sha256': sha(target)}


if __name__ == '__main__':
    action = sys.argv[1]
    spec = Q4_SPEC if '--q4' in sys.argv else DEFAULT_SPEC
    if action == 'prepare':
        prepare(spec)
    else:
        results = []
        for n in (2,3):
            if action == 'finalize': preserve(n, spec)
            results.append(verify(n, spec=spec))
        write_json(spec.out/'readback.json', results)
        print(json.dumps(results,ensure_ascii=False))
