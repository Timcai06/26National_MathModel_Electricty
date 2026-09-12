"""Read C-problem source structure without cleaning data or editing workbooks.

Run from the project root: rtk proxy python3 scripts/inspect_source_headers.py
Requires openpyxl. Output is structural evidence, not a full data audit.
"""

import hashlib
import json
import platform
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "CUMCM2026Problems/C题"
OUTPUT = ROOT / "analysis/source_inspection.json"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inspect_workbook(path):
    before = digest(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheets = []
    try:
        for sheet in book:
            samples = {}
            # First rows, last rows, and first-column labels expose source schemas
            # and abbreviated template date blocks without interpreting timestamps.
            selected_rows = set(range(1, min(6, sheet.max_row) + 1))
            selected_rows.update(range(max(1, sheet.max_row - 3), sheet.max_row + 1))
            for row in sheet.iter_rows():
                row_number = next((c.row for c in row if c.value is not None), None)
                if row_number is None:
                    continue
                for cell in row:
                    if cell.value is not None and (
                        row_number in selected_rows or cell.column == 1
                    ):
                        samples[cell.coordinate] = {
                            "value": cell.value,
                            "data_type": cell.data_type,
                            "number_format": cell.number_format,
                        }
            sheets.append({
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "samples": samples,
            })
    finally:
        book.close()
    after = digest(path)
    if before != after:
        raise RuntimeError(f"Source changed during inspection: {path}")
    return {
        "path": str(path.relative_to(ROOT)),
        "sha256_before": before,
        "sha256_after": after,
        "sheets": sheets,
    }


def main():
    pdf = SOURCE / "C题.pdf"
    report = {
        "inspected_at_utc": datetime.now(timezone.utc).isoformat(),
        "scope": "source dimensions and cell samples; no quality audit or optimization",
        "python": platform.python_version(),
        "openpyxl": openpyxl.__version__,
        "pdf": {"path": str(pdf.relative_to(ROOT)), "sha256": digest(pdf)},
        "workbooks": [inspect_workbook(p) for p in sorted(SOURCE.rglob("*.xlsx"))],
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n")
    print(f"Inspected {len(report['workbooks'])} workbooks; source hashes unchanged.")
    print(OUTPUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
