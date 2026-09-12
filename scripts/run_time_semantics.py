"""Settle the input/output time semantics from the official files themselves.

Read-only on every source. Emits an evidence bundle and the integration-rule
sensitivity used to freeze the paper's time contract. Runs no annual model.
"""
import csv, hashlib, json, sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import numpy as np
import openpyxl

from src.core import BatteryParams
from src.q1.compare import compare_q1_strategies

PROBLEM = ROOT / 'CUMCM2026Problems/C题'
TEMPLATES = ['result1.xlsx', 'result2.xlsx', 'result3.xlsx', 'result4-2.xlsx', 'result4-3.xlsx']
OUT = ROOT / 'analysis/time_semantics_v1'


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def cell_text(value):
    if value is None:
        return None
    if hasattr(value, 'strftime') and not isinstance(value, str):
        return value.strftime('%H:%M') if not hasattr(value, 'date') else value.strftime('%Y-%m-%d')
    return str(value)


def read_evidence():
    """Collect the label evidence from attachment 1 and all five templates."""
    hashes = {}
    a1 = PROBLEM / '附件/附件1.xlsx'
    hashes[str(a1.relative_to(ROOT))] = sha256(a1)
    ws = openpyxl.load_workbook(a1, data_only=True)['Sheet1']
    input_labels = [cell_text(ws.cell(r, 1).value) for r in range(2, 146)]
    if len(input_labels) != 144:
        raise ValueError('Attachment 1 must carry exactly 144 sample labels')

    templates = {}
    for name in TEMPLATES:
        path = PROBLEM / '附件/附件5' / name
        hashes[str(path.relative_to(ROOT))] = sha256(path)
        wb = openpyxl.load_workbook(path)
        entry = {'sheets': wb.sheetnames}
        # Purchase-axis labels: column header row for annual sheets, column A for result1.
        plan = wb['计划购电量']
        if name == 'result1.xlsx':
            axis = [cell_text(plan.cell(r, 1).value) for r in range(2, 146)]
        else:
            axis = [cell_text(plan.cell(1, c).value) for c in range(2, 146)]
        entry['purchase_axis_first'] = axis[0]
        entry['purchase_axis_last'] = axis[-1]
        entry['purchase_axis_n'] = len(axis)
        # Battery sheet: the four-hour buckets and the two SOC instants.
        bat = wb['充放电量']
        col = 1 if name == 'result1.xlsx' else 2
        entry['battery_buckets'] = [cell_text(bat.cell(r, col).value) for r in range(2, 8)]
        inst_col = col + 3
        entry['battery_instants'] = [cell_text(bat.cell(r, inst_col).value) for r in range(2, 4)]
        templates[name] = entry
        if sha256(path) != hashes[str(path.relative_to(ROOT))]:
            raise RuntimeError('Template changed while being read')
    if sha256(a1) != hashes[str(a1.relative_to(ROOT))]:
        raise RuntimeError('Attachment 1 changed while being read')
    return input_labels, templates, hashes


def check_contradiction(input_labels, templates):
    """The templates assert two incompatible day windows. Record both literally."""
    buckets = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00', '16:00-20:00', '20:00-24:00']
    findings = {}
    for name, entry in templates.items():
        if entry['battery_buckets'] != buckets:
            raise ValueError(f'{name}: unexpected battery buckets {entry["battery_buckets"]}')
        instants = [x[1:] if isinstance(x, str) and x.startswith('00:') else x
                    for x in entry['battery_instants']]
        if instants != ['0:00', '24:00']:
            raise ValueError(f'{name}: unexpected SOC instants {entry["battery_instants"]}')
        findings[name] = {
            'battery_sheet_window': '00:00-24:00',
            'purchase_axis_window': f"{entry['purchase_axis_first']} .. {entry['purchase_axis_last']}",
            'consistent': entry['purchase_axis_first'] == '0:00-0:10',
        }
    return {
        'input_first_label': input_labels[0],
        'input_last_label': input_labels[-1],
        'battery_buckets': buckets,
        'battery_instants': ['0:00', '24:00'],
        'per_template': findings,
        'conclusion': (
            'Every template battery sheet tiles 00:00-24:00 and pins SOC at 0:00 and 24:00, '
            'matching the problem statement. Every purchase axis instead starts at 0:10-0:20, '
            'i.e. one position later. The two sheets of the same official file cannot both be '
            'literal, so the purchase-axis label text is the inconsistent one.'
        ),
    }


def integration_rules(power_prev_end, power):
    """Right-hold, trapezoid and left-hold energies for one day of 144 samples.

    power[k] is the instantaneous sample at the END of slot k. power_prev_end is
    the sample at 00:00, i.e. the previous day's 24:00 reading.
    """
    power = np.asarray(power, dtype=float)
    boundaries = np.r_[float(power_prev_end), power]
    return {
        'right_hold': power / 6.0,
        'trapezoid': (boundaries[:-1] + boundaries[1:]) / 12.0,
        'left_hold': boundaries[:-1] / 6.0,
    }


def annual_integration_gap():
    """Quantify the rule choice on the real year; the daily gap is analytic."""
    source = ROOT / 'data/processed/team_delivery/annual_actual.csv'
    with source.open(newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    days = defaultdict(list)
    for row in rows:
        days[row['date']].append(row)
    ordered = sorted(days)
    prev_end = {'load': None, 'pv': None}
    per_day, slot_gaps = [], {'load': [], 'pv': []}
    for day in ordered:
        block = sorted(days[day], key=lambda r: int(r['slot']))
        if [int(r['slot']) for r in block] != list(range(144)):
            raise ValueError(f'{day}: incomplete or unordered slots; no repair')
        record = {'date': day}
        for key, column in (('load', 'load_kw'), ('pv', 'pv_actual_kw')):
            power = np.array([float(r[column]) for r in block])
            start = power[-1] if prev_end[key] is None else prev_end[key]
            first_day = prev_end[key] is None
            rules = integration_rules(start, power)
            gap = rules['right_hold'].sum() - rules['trapezoid'].sum()
            analytic = (power[-1] - start) / 12.0
            if abs(gap - analytic) > 1e-9:
                raise ValueError('Daily gap identity violated')
            record[f'{key}_right_hold_kwh'] = float(rules['right_hold'].sum())
            record[f'{key}_trapezoid_kwh'] = float(rules['trapezoid'].sum())
            record[f'{key}_daily_gap_kwh'] = float(gap)
            record[f'{key}_first_day_cyclic_start'] = first_day
            slot_gaps[key].append(np.abs(rules['right_hold'] - rules['trapezoid']).max())
            prev_end[key] = power[-1]
        per_day.append(record)
    summary = {}
    for key in ('load', 'pv'):
        right = sum(r[f'{key}_right_hold_kwh'] for r in per_day)
        trap = sum(r[f'{key}_trapezoid_kwh'] for r in per_day)
        # Consecutive days share a midnight, so the daily gaps telescope to the
        # two ends of the year. The first day starts from its own 24:00 sample.
        column = 'load_kw' if key == 'load' else 'pv_actual_kw'
        ends = [float(sorted(days[d], key=lambda r: int(r['slot']))[-1][column]) for d in (ordered[0], ordered[-1])]
        telescoped = (ends[1] - ends[0]) / 12.0
        if abs((right - trap) - telescoped) > 1e-6:
            raise ValueError(f'{key}: annual gap does not telescope to the year ends')
        summary[key] = {
            'annual_right_hold_kwh': right,
            'annual_trapezoid_kwh': trap,
            'annual_gap_kwh': right - trap,
            'annual_gap_relative': (right - trap) / right if right else 0.0,
            'max_abs_daily_gap_kwh': max(abs(r[f'{key}_daily_gap_kwh']) for r in per_day),
            'max_abs_slot_gap_kwh': float(max(slot_gaps[key])),
            'telescoped_gap_kwh': telescoped,
            'telescoped_from': {'first_day_2400_kw': ends[0], 'last_day_2400_kw': ends[1]},
        }
    return {
        'source': str(source.relative_to(ROOT)),
        'source_sha256': sha256(source),
        'days': len(per_day),
        'identity': 'daily right-hold total minus trapezoid total equals (P(24:00) - P(00:00)) / 12',
        'annual_identity': 'consecutive days share a midnight, so the daily gaps telescope: the annual gap '
                           'equals (last day 24:00 sample - first day 00:00 boundary) / 12',
        'first_day_start_rule': '2025-01-01 has no previous day; its 00:00 boundary uses its own 24:00 sample',
        'summary': summary,
        'per_day': per_day,
    }


def q1_under_each_rule():
    """Re-solve the representative day under each integration rule."""
    source = ROOT / 'data/processed/team_delivery/day_base.csv'
    with source.open(newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    if [int(r['slot']) for r in rows] != list(range(144)):
        raise ValueError('day_base slots must already be ordered 0..143')
    load_kw = np.array([float(r['load_kw']) for r in rows])
    pv_kw = np.array([float(r['pv_forecast_kw']) for r in rows])
    price = np.array([float(r['price']) for r in rows])
    # Q1 states every day is identical, so the 00:00 boundary is the same day's 24:00.
    load = integration_rules(load_kw[-1], load_kw)
    pv = integration_rules(pv_kw[-1], pv_kw)
    price_trapezoid = (np.r_[price[-1], price][:-1] + np.r_[price[-1], price][1:]) / 2.0
    cases = {
        'right_hold': (load['right_hold'], pv['right_hold'], price),
        'trapezoid': (load['trapezoid'], pv['trapezoid'], price),
        'left_hold': (load['left_hold'], pv['left_hold'], price),
        'trapezoid_with_averaged_price': (load['trapezoid'], pv['trapezoid'], price_trapezoid),
    }
    results = {}
    for name, (l, v, p) in cases.items():
        report = compare_q1_strategies(
            load_kwh=l, pv_forecast_kwh=v, price_yuan_per_kwh=p,
            initial_energy_kwh=6000, battery_params=BatteryParams(.9, .9, 'bus'),
            slot_ids=list(range(144)),
            input_metadata=dict(input_kind='standardized', data_version='sha256:' + sha256(source),
                                mapping_version='interval-end-v1', assumptions_version=f'integration-{name}'),
        )
        if report['status'] != 'passed':
            raise RuntimeError(f'{name}: Q1 comparison failed')
        results[name] = {row['strategy']: {'cost_yuan': row['cost_yuan'],
                                           'purchase_kwh': row['purchase_kwh'],
                                           'saving_fraction_vs_no_storage': row['saving_fraction_vs_no_storage']}
                         for row in report['summary']}
        results[name]['load_total_kwh'] = float(l.sum())
        results[name]['pv_total_kwh'] = float(v.sum())
    base = results['right_hold']['optimized']['cost_yuan']
    for name, entry in results.items():
        entry['optimized_delta_vs_right_hold_yuan'] = entry['optimized']['cost_yuan'] - base
    return {'source': str(source.relative_to(ROOT)), 'source_sha256': sha256(source), 'cases': results}


def paper_table_slots():
    """The six 表1 intervals and the six 表2 buckets under the frozen contract."""
    rows = []
    for hour in (10, 12, 14, 16, 18, 20):
        slot = hour * 6
        sample = (datetime(2000, 1, 1) + timedelta(minutes=10 * (slot + 1))).strftime('%H:%M')
        column = slot + 2
        letters = ''
        n = column
        while n:
            n, r = divmod(n - 1, 26)
            letters = chr(65 + r) + letters
        rows.append({'paper_interval': f'{hour}:00-{hour}:10', 'slot_id': slot,
                     'attachment_sample_label': sample, 'template_column': letters,
                     'template_column_printed_label': f'{hour}:10-{hour}:20'})
    buckets = [{'bucket': f'{h}:00-{h + 4}:00', 'slot_ids': f'{h * 6}..{h * 6 + 23}'} for h in range(0, 24, 4)]
    return {'table1_intervals': rows, 'table2_buckets': buckets,
            'note': 'Values are filled positionally. The printed purchase-axis label is one position later.'}


def main():
    input_labels, templates, hashes = read_evidence()
    bundle = {
        'decision': 'interval-end-v1',
        'scope': 'Time-semantics freeze for the paper and the five result workbooks. No annual model was re-run.',
        'statement': {
            'input_sample': 'A value labelled T in attachments 1, 2 and 4 is the instantaneous sample at T and is '
                            'assigned to the slot ENDING at T, i.e. slot k covers [k*10min, (k+1)*10min).',
            'model_day': 'Every problem runs over 00:00-24:00 with 144 slots and 145 state boundaries.',
            'energy_rule': 'Right-endpoint hold: E_k = P(t_{k+1}) / 6. Declared as an approximation, not as a claim '
                           'that the attachments provide interval-average power.',
            'template_fill': 'Positional. Slot k is written to the k-th purchase column. The printed purchase-axis '
                             'label of that column reads one position later; the official template is left unmodified.',
            'forecast_rule': 'Attachment 3 hourly forecasts are piecewise-linear in time and integrated exactly onto '
                             'the same slots, so plan and settlement share one slot grid.',
        },
        'source_sha256': hashes,
        'label_evidence': check_contradiction(input_labels, templates),
        'integration_sensitivity': annual_integration_gap(),
        'q1_under_each_rule': q1_under_each_rule(),
        'paper_table_mapping': paper_table_slots(),
    }
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'result.json').write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    per_day = bundle['integration_sensitivity'].pop('per_day')
    with (OUT / 'daily_integration_gap.csv').open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=list(per_day[0]))
        writer.writeheader()
        writer.writerows(per_day)
    (OUT / 'contract.json').write_text(json.dumps(
        {'decision': bundle['decision'], 'statement': bundle['statement'],
         'source_sha256': hashes}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'decision': bundle['decision'],
                      'label_evidence': bundle['label_evidence']['conclusion'],
                      'integration': bundle['integration_sensitivity']['summary'],
                      'q1': {k: v.get('optimized_delta_vs_right_hold_yuan')
                             for k, v in bundle['q1_under_each_rule']['cases'].items()}},
                     ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
