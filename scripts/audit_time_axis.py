"""Audit all C-problem 10-minute axes; export a candidate, never fill templates."""

import csv
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.time_axis import build_day_slots, minute_offset

SOURCE = ROOT / "CUMCM2026Problems/C题/附件"
OUTPUT = ROOT / "analysis/time_axis"


def checksum(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def axis(book, sheet_name, coordinates):
    sheet = book[sheet_name]
    return [{"cell": coord, "value": str(sheet[coord].value),
             "number_format": sheet[coord].number_format}
            for coord in coordinates]


def clock(minutes):
    day, within_day = divmod(minutes, 1440)
    hours, minute = divmod(within_day, 60)
    return f"{hours}:{minute:02d}" + (f"+{day}" if day else "")


def main():
    files = sorted(SOURCE.rglob("*.xlsx"))
    before = {str(p.relative_to(ROOT)): checksum(p) for p in files}
    input_axes, output_axes = {}, {}
    columns = [get_column_letter(c) for c in range(2, 146)]
    for path in files:
        # Normal read mode allows bounded random header access efficiently.
        # There is deliberately no save call anywhere in this script.
        book = openpyxl.load_workbook(path, data_only=False)
        try:
            name = path.name
            if name == "附件1.xlsx":
                if book.active.max_row != 145:
                    raise ValueError("Unexpected attachment 1 row count")
                input_axes[name + "/Sheet1"] = axis(book, "Sheet1", [f"A{r}" for r in range(2, 146)])
            elif name in {"附件2.xlsx", "附件4.xlsx"}:
                for sheet in book:
                    if sheet.max_column != 145:
                        raise ValueError(f"Unexpected input column count: {name}/{sheet.title}")
                    input_axes[name + "/" + sheet.title] = axis(book, sheet.title, [c + "1" for c in columns])
            elif name.startswith("result"):
                for sheet_name in ["计划购电量", "调整购电量"]:
                    if sheet_name in book.sheetnames:
                        coords = ([f"A{r}" for r in range(2, 146)] if name == "result1.xlsx"
                                  else [c + "1" for c in columns])
                        output_axes[name + "/" + sheet_name] = axis(book, sheet_name, coords)
        finally:
            book.close()

    if len(input_axes) != 4 or len(output_axes) != 7:
        raise ValueError("Expected four input axes and seven output axes")
    expected = list(range(10, 1441, 10))
    for name, cells in input_axes.items():
        if [minute_offset(c["value"]) for c in cells] != expected:
            raise ValueError(f"Input label mismatch: {name}")

    discrepancies, formatting_variants = [], []
    for name, cells in output_axes.items():
        for i, cell in enumerate(cells):
            raw = cell["value"]
            if raw == "7:0-7:10":
                formatting_variants.append({"axis": name, "cell": cell["cell"], "raw": raw})
            start, end = raw.split("-")
            parsed = (minute_offset(start), minute_offset(end))
            desired = ((i + 1) * 10, (i + 2) * 10)
            if parsed != desired:
                # This exception is observed source text, not a corrected input.
                if not (i == 143 and not name.startswith("result1.xlsx/")
                        and raw == "0:00-0:10+1"):
                    raise ValueError(f"Unexpected output label: {name}/{cell['cell']}={raw}")
                discrepancies.append({"axis": name, "cell": cell["cell"], "raw": raw,
                                      "literal_offsets": parsed,
                                      "sequence_based_candidate_offsets": desired})

    labels = [c["value"] for c in input_axes["附件1.xlsx/Sheet1"]]
    # Anchor date is illustrative; the exported contract uses day-relative offsets.
    slots = build_day_slots(date(2025, 1, 1), labels, semantics="interval_end")
    q1 = output_axes["result1.xlsx/计划购电量"]
    annual = output_axes["result2.xlsx/计划购电量"]
    rows = []
    for i, slot in enumerate(slots):
        rows.append({
            "mapping_version": "interval-end-v1",
            "slot_id": i,
            "model_start_minute": i * 10,
            "model_end_minute": (i + 1) * 10,
            "model_interval": clock(i * 10) + "-" + clock((i + 1) * 10),
            "input_label_value": labels[i],
            "attachment1_time_cell": f"A{i + 2}",
            "attachment1_value_row": i + 2,
            "annual_input_column": columns[i],
            "q1_output_label_cell": q1[i]["cell"],
            "q1_output_value_cell": f"B{i + 2}",
            "q1_output_label_original": q1[i]["value"],
            "annual_output_column": columns[i],
            "annual_output_label_original": annual[i]["value"],
            "output_display_offset_minutes": 10,
            "last_label_ambiguity": i == 143,
            "state_start_index": i,
            "state_end_index": i + 1,
        })
    after = {str(p.relative_to(ROOT)): checksum(p) for p in files}
    if before != after:
        raise RuntimeError("Source files changed during audit")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    with (OUTPUT / "slot_mapping_candidate.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    report = {
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": "structural_checks_passed_semantics_candidate",
        "scope": "four input axes and seven purchase/adjustment output axes; not full Gate 0",
        "input_axes": input_axes, "output_axes": output_axes,
        "input_cells_checked": 4 * 144, "output_cells_checked": 7 * 144,
        "output_label_ambiguities": discrepancies,
        "output_formatting_variants": formatting_variants,
        "candidate_interpretation": "interval_end; output stays in original positional cell",
        "unresolved": ["physical interval versus official output label conflict",
                       "input point values treated as interval representative power",
                       "attachment 3 interpolation and full data validation",
                       "independent A/B crosscheck and final output semantic acceptance"],
        "source_sha256_before": before, "source_sha256_after": after,
    }
    (OUTPUT / "audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(f"PASS: {4 * 144} input labels, {7 * 144} output labels; {len(discrepancies)} annual last-label ambiguities.")
    print("144 candidate mapping rows exported; all 9 source workbook hashes unchanged.")
    print("Physical/output semantic conflict remains explicit; Gate 0 is not passed.")


if __name__ == "__main__":
    main()
